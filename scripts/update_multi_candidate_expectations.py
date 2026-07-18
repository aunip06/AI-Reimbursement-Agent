from __future__ import annotations

from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook


WORKBOOK_PATHS = [
    Path("input/AI_Reimbursement_Smoke_Test.xlsx"),
    Path("input/AI_Reimbursement_Comprehensive_Test.xlsx"),
]

WORKSHEET_NAME = "2026-05_May"

EXPECTATION_UPDATES = {
    "TEST001_MAY_007A": {
        "Expected_Status": "DUPLICATE_RECEIPT_REFERENCE",
        "Remarks": (
            "First claim selects the same isolated receipt "
            "candidate as TEST001_MAY_007B."
        ),
        "Test_Purpose": (
            "Tests duplicate protection using the selected "
            "receipt fingerprint rather than only the PDF page."
        ),
    },
    "TEST001_MAY_007B": {
        "Expected_Status": "DUPLICATE_RECEIPT_REFERENCE",
        "Remarks": (
            "Second claim selects the same isolated receipt "
            "candidate as TEST001_MAY_007A."
        ),
        "Test_Purpose": (
            "Tests duplicate protection using the selected "
            "receipt fingerprint rather than only the PDF page."
        ),
    },
    "TEST001_MAY_008": {
        "Expected_Status": "OK_AUTO_APPROVED",
        "Remarks": (
            "The page contains multiple separate screenshots. "
            "The claim uniquely matches one isolated candidate."
        ),
        "Test_Purpose": (
            "Tests independent screenshot segmentation, OCR and "
            "unique claim-to-candidate matching."
        ),
    },
}


def update_workbook(workbook_path: Path) -> None:
    if not workbook_path.exists():
        raise FileNotFoundError(
            f"Workbook not found: {workbook_path}"
        )

    backup_path = workbook_path.with_name(
        f"{workbook_path.stem}"
        f".pre_multi_candidate_backup"
        f"{workbook_path.suffix}"
    )

    if not backup_path.exists():
        copy2(
            workbook_path,
            backup_path,
        )

        print(
            f"Backup created: {backup_path}"
        )

    workbook = load_workbook(
        workbook_path
    )

    if WORKSHEET_NAME not in workbook.sheetnames:
        raise ValueError(
            f"Worksheet '{WORKSHEET_NAME}' was not found "
            f"in {workbook_path}."
        )

    worksheet = workbook[
        WORKSHEET_NAME
    ]

    headers = {
        str(cell.value).strip(): cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }

    required_headers = {
        "Expense_ID",
        "Expected_Status",
        "Remarks",
        "Test_Purpose",
    }

    missing_headers = (
        required_headers
        - set(headers)
    )

    if missing_headers:
        raise ValueError(
            f"Missing columns in {workbook_path}: "
            f"{sorted(missing_headers)}"
        )

    updated_expense_ids: set[str] = set()

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        expense_id_value = worksheet.cell(
            row=row_number,
            column=headers["Expense_ID"],
        ).value

        if expense_id_value is None:
            continue

        expense_id = str(
            expense_id_value
        ).strip()

        update_values = (
            EXPECTATION_UPDATES.get(
                expense_id
            )
        )

        if update_values is None:
            continue

        for column_name, value in update_values.items():
            worksheet.cell(
                row=row_number,
                column=headers[column_name],
                value=value,
            )

        updated_expense_ids.add(
            expense_id
        )

    missing_expense_ids = (
        set(EXPECTATION_UPDATES)
        - updated_expense_ids
    )

    if missing_expense_ids:
        raise ValueError(
            f"Expected test rows were not found in "
            f"{workbook_path}: "
            f"{sorted(missing_expense_ids)}"
        )

    workbook.save(
        workbook_path
    )

    print(
        f"Updated: {workbook_path}"
    )

    for expense_id in sorted(
        updated_expense_ids
    ):
        print(
            f"  {expense_id}: "
            f"{EXPECTATION_UPDATES[expense_id]['Expected_Status']}"
        )


def main() -> None:
    for workbook_path in WORKBOOK_PATHS:
        update_workbook(
            workbook_path
        )

    print(
        "\nMulti-candidate test expectations updated successfully."
    )


if __name__ == "__main__":
    main()