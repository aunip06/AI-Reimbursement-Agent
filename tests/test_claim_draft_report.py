"""Regression tests for PDF-first claim-draft workbook generation."""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from tools.claim_draft_report_tool import (
    validate_completed_claim_draft,
    write_claim_draft_workbook,
)


def receipt(**overrides):
    """Build a minimal receipt-like object for deterministic tests."""

    values = {
        "bill_found": True,
        "receipt_type": "upi_screenshot",
        "date": "2026-04-16",
        "amount": 255.0,
        "currency": "INR",
        "merchant": "ANAND HELODE",
        "payment_mode": "upi",
        "payment_app": "PhonePe",
        "message": "Transaction Successful",
        "remark": "Flyer Printouts",
        "transaction_id": "T2604160917495671535032",
        "utr": "034410846237",
        "invoice_no": None,
        "receipt_no": None,
        "gst_number": None,
        "confidence": "high",
        "extraction_notes": "Clear payment screenshot.",
    }
    values.update(overrides)
    return SimpleNamespace(**values)


def intake_item(
    generated_id: str,
    *,
    receipt_value=None,
    transaction_id: str | None = None,
):
    """Build a minimal PDFIntakeItem-like object."""

    receipt_object = (
        receipt(
            transaction_id=transaction_id
        )
        if receipt_value is None
        else receipt_value
    )

    return SimpleNamespace(
        generated_id=generated_id,
        source_file_name="exp_apr.pdf",
        page_number=1,
        candidate_no=1,
        status="EXTRACTED_READY",
        result_source="openai",
        receipt=receipt_object,
        candidate_image_path="candidate.png",
        candidate_map_image_path="map.png",
        ocr_text="Transaction Successful",
        candidate_fingerprint=generated_id,
        perceptual_hash=generated_id[-8:],
        duplicate_of_generated_id=None,
        message="Extracted.",
    )


def intake_result(items):
    """Build a minimal PDFIntakeResult-like object."""

    ingestion = SimpleNamespace(
        source_file_name="exp_apr.pdf",
        source_family="pdf",
        status="READY",
        converter="pdf_passthrough",
        document_fingerprint="doc-fingerprint",
        model_dump=lambda mode="json": {
            "source_file_name": "exp_apr.pdf",
            "source_family": "pdf",
            "status": "READY",
            "converter": "pdf_passthrough",
            "document_fingerprint": "doc-fingerprint",
        },
    )

    return SimpleNamespace(
        ingestion=ingestion,
        total_pages=2,
        total_candidates=len(items),
        extracted_count=len(items),
        exception_count=0,
        duplicate_count=0,
        items=items,
    )


def test_generate_and_validate_claim_draft(tmp_path: Path) -> None:
    """One receipt becomes one draft row and validates after confirmation."""

    source = tmp_path / "claim_draft.xlsx"
    result = intake_result(
        [
            intake_item(
                "APR-P01-C01",
                transaction_id="TXN-APR-001",
            )
        ]
    )

    write_claim_draft_workbook(
        results=[result],
        output_path=source,
    )

    workbook = load_workbook(source)
    assert "Claim_Draft" in workbook.sheetnames
    assert "Evidence_Register" in workbook.sheetnames
    assert "Needs_Attention" in workbook.sheetnames
    assert workbook["Claim_Draft"].max_row == 2

    claim_sheet = workbook["Claim_Draft"]
    headers = {
        cell.value: cell.column
        for cell in claim_sheet[1]
    }
    row_number = 2

    claim_sheet.cell(
        row=row_number,
        column=headers["Employee_Name"],
        value="Demo Employee",
    )
    claim_sheet.cell(
        row=row_number,
        column=headers["Expense_Category"],
        value="Printing",
    )
    claim_sheet.cell(
        row=row_number,
        column=headers["Business_Purpose"],
        value="Official flyer printing",
    )
    claim_sheet.cell(
        row=row_number,
        column=headers["User_Confirmation"],
        value="Confirmed",
    )
    claim_sheet.cell(
        row=row_number,
        column=headers["Manager_Approval"],
        value="Approved",
    )
    workbook.save(source)

    validated_path, rows = validate_completed_claim_draft(
        input_workbook=source,
    )

    assert validated_path.exists()
    assert rows[0]["Validation_Status"] == "OK_VERIFIED"


def test_cross_document_duplicate_is_blocked(tmp_path: Path) -> None:
    """Repeated transaction identifiers are not treated as new claims."""

    source = tmp_path / "duplicates.xlsx"
    first = intake_result(
        [
            intake_item(
                "APR-P01-C01",
                transaction_id="SAME-TXN-001",
            )
        ]
    )
    second_item = intake_item(
        "MAY-P01-C01",
        transaction_id="SAME-TXN-001",
    )
    second_item.source_file_name = "exp_may.pdf"
    second = intake_result(
        [second_item]
    )
    second.ingestion.source_file_name = "exp_may.pdf"

    write_claim_draft_workbook(
        results=[first, second],
        output_path=source,
    )

    workbook = load_workbook(source)
    claim_sheet = workbook["Claim_Draft"]
    headers = {
        cell.value: cell.column
        for cell in claim_sheet[1]
    }
    statuses = [
        claim_sheet.cell(
            row=row,
            column=headers["Draft_Status"],
        ).value
        for row in range(
            2,
            claim_sheet.max_row + 1,
        )
    ]

    assert "DUPLICATE_RECEIPT" in statuses
