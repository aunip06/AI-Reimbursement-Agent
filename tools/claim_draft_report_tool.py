"""Create and validate PDF-first reimbursement claim-draft workbooks.

This module connects the existing document-intake pipeline to a controlled
PDF -> evidence register -> editable claim draft -> validation workflow.

Safety boundary:
- Evidence_Register is machine-generated from isolated receipt candidates.
- Claim_Draft contains user-editable business fields.
- Validation always compares Claim_Draft values against Evidence_Register.
- This feature does not approve or pay a reimbursement.
"""

from __future__ import annotations

import re
from collections import defaultdict
from copy import copy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from models.pdf_intake import PDFIntakeItem, PDFIntakeResult


CLAIM_DRAFT_HEADERS = [
    "No",
    "Receipt_ID",
    "Month",
    "Extracted_Date",
    "Claimed_Date",
    "Extracted_Amount",
    "Claimed_Amount",
    "Currency",
    "Merchant_Payee",
    "Payment_Mode",
    "Transaction_ID",
    "UTR",
    "Extracted_Remark",
    "Expense_Category",
    "Business_Purpose",
    "Employee_ID",
    "Employee_Name",
    "Project_Cost_Centre",
    "User_Confirmation",
    "Manager_Approval",
    "Draft_Status",
    "Source_PDF",
    "Page_No",
    "Candidate_No",
]

EVIDENCE_HEADERS = [
    "Receipt_ID",
    "Source_PDF",
    "Input_Family",
    "Document_Fingerprint",
    "Page_No",
    "Candidate_No",
    "Intake_Status",
    "Result_Source",
    "Bill_Found",
    "Receipt_Type",
    "Extracted_Date",
    "Extracted_Amount",
    "Currency",
    "Merchant_Payee",
    "Payment_Mode",
    "Payment_App",
    "Payment_Message",
    "Extracted_Remark",
    "Transaction_ID",
    "UTR",
    "Invoice_No",
    "Receipt_No",
    "GST_Number",
    "Extraction_Confidence",
    "Extraction_Notes",
    "Duplicate_Of",
    "Cross_Document_Duplicate_Of",
    "Candidate_Fingerprint",
    "Perceptual_Hash",
    "Evidence_Image_Path",
    "Candidate_Map_Path",
    "OCR_Text",
    "System_Message",
]

VALIDATION_HEADERS = [
    "Receipt_ID",
    "Employee_ID",
    "Employee_Name",
    "Expense_Category",
    "Business_Purpose",
    "Evidence_Date",
    "Claimed_Date",
    "Evidence_Amount",
    "Claimed_Amount",
    "User_Confirmation",
    "Manager_Approval",
    "Validation_Status",
    "Validation_Reason",
    "Source_PDF",
    "Page_No",
    "Candidate_No",
]

CATEGORY_VALUES = [
    "Travel",
    "Food",
    "Hotel",
    "Printing",
    "Stationery",
    "Fuel",
    "Material",
    "Client Meeting",
    "Training",
    "Other",
]

CONFIRMATION_VALUES = [
    "Pending",
    "Confirmed",
    "Rejected",
]

APPROVAL_VALUES = [
    "Pending",
    "Approved",
    "Rejected",
]

HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"
EXTRACTED_FILL = "D9EAF7"
EDITABLE_FILL = "FFF2CC"
GOOD_FILL = "C6EFCE"
WARNING_FILL = "FCE4D6"
ERROR_FILL = "FFC7CE"
GREY_FILL = "E7E6E6"


def normalise_identifier(value: str | None) -> str:
    """Return a comparison-safe transaction or invoice identifier."""

    if not value:
        return ""

    return re.sub(
        r"[^a-z0-9]",
        "",
        str(value).casefold(),
    )


def normalise_amount(value: Any) -> Decimal | None:
    """Convert a spreadsheet or extracted amount into two-decimal Decimal."""

    if value is None or value == "":
        return None

    try:
        cleaned = (
            str(value)
            .replace("₹", "")
            .replace(",", "")
            .strip()
        )
        return Decimal(cleaned).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def normalise_date(value: Any) -> str | None:
    """Convert supported date values into ISO YYYY-MM-DD text."""

    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, date):
        return value.isoformat()

    text = str(value).strip()

    for format_string in (
        "%Y-%m-%d",
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%d %b %Y",
        "%d %B %Y",
    ):
        try:
            return datetime.strptime(
                text,
                format_string,
            ).date().isoformat()
        except ValueError:
            continue

    return None


def month_label(value: Any) -> str:
    """Return Month-YYYY, or Unknown Month when a date is unavailable."""

    iso_date = normalise_date(value)

    if not iso_date:
        return "Unknown Month"

    parsed = datetime.strptime(
        iso_date,
        "%Y-%m-%d",
    )

    return parsed.strftime("%B-%Y")


def item_receipt(item: PDFIntakeItem):
    """Return an item's receipt model, if available."""

    return item.receipt


def duplicate_key(item: PDFIntakeItem) -> str | None:
    """Create a cross-document duplicate key from strong identifiers."""

    receipt = item_receipt(item)

    if receipt is None or not receipt.bill_found:
        return None

    for name, value in (
        ("transaction_id", receipt.transaction_id),
        ("utr", receipt.utr),
        ("invoice_no", receipt.invoice_no),
        ("receipt_no", receipt.receipt_no),
    ):
        normalised = normalise_identifier(value)

        if len(normalised) >= 6:
            return f"{name}:{normalised}"

    if item.perceptual_hash:
        return f"visual:{item.perceptual_hash}"

    if item.candidate_fingerprint:
        return f"exact:{item.candidate_fingerprint}"

    return None


def cross_document_duplicate_map(
    results: Iterable[PDFIntakeResult],
) -> dict[str, str]:
    """Map later duplicate receipt IDs to the first receipt ID."""

    grouped: dict[str, list[PDFIntakeItem]] = defaultdict(list)

    for result in results:
        for item in result.items:
            key = duplicate_key(item)

            if key:
                grouped[key].append(item)

    duplicate_of: dict[str, str] = {}

    for items in grouped.values():
        ordered = sorted(
            items,
            key=lambda item: (
                item.source_file_name.casefold(),
                item.page_number,
                item.candidate_no or 0,
                item.generated_id,
            ),
        )

        if len(ordered) < 2:
            continue

        original_id = ordered[0].generated_id

        for duplicate_item in ordered[1:]:
            duplicate_of[duplicate_item.generated_id] = original_id

    return duplicate_of


def evidence_row(
    result: PDFIntakeResult,
    item: PDFIntakeItem,
    duplicate_of: dict[str, str],
) -> dict[str, Any]:
    """Convert one intake item into a complete locked evidence row."""

    receipt = item_receipt(item)
    ingestion = result.ingestion

    return {
        "Receipt_ID": item.generated_id,
        "Source_PDF": item.source_file_name,
        "Input_Family": ingestion.source_family,
        "Document_Fingerprint": ingestion.document_fingerprint,
        "Page_No": item.page_number,
        "Candidate_No": item.candidate_no,
        "Intake_Status": item.status,
        "Result_Source": item.result_source,
        "Bill_Found": receipt.bill_found if receipt else False,
        "Receipt_Type": receipt.receipt_type if receipt else None,
        "Extracted_Date": receipt.date if receipt else None,
        "Extracted_Amount": receipt.amount if receipt else None,
        "Currency": receipt.currency if receipt else None,
        "Merchant_Payee": receipt.merchant if receipt else None,
        "Payment_Mode": receipt.payment_mode if receipt else None,
        "Payment_App": receipt.payment_app if receipt else None,
        "Payment_Message": receipt.message if receipt else None,
        "Extracted_Remark": receipt.remark if receipt else None,
        "Transaction_ID": receipt.transaction_id if receipt else None,
        "UTR": receipt.utr if receipt else None,
        "Invoice_No": receipt.invoice_no if receipt else None,
        "Receipt_No": receipt.receipt_no if receipt else None,
        "GST_Number": receipt.gst_number if receipt else None,
        "Extraction_Confidence": receipt.confidence if receipt else None,
        "Extraction_Notes": receipt.extraction_notes if receipt else None,
        "Duplicate_Of": item.duplicate_of_generated_id,
        "Cross_Document_Duplicate_Of": duplicate_of.get(item.generated_id),
        "Candidate_Fingerprint": item.candidate_fingerprint,
        "Perceptual_Hash": item.perceptual_hash,
        "Evidence_Image_Path": item.candidate_image_path,
        "Candidate_Map_Path": item.candidate_map_image_path,
        "OCR_Text": item.ocr_text,
        "System_Message": item.message,
    }


def initial_draft_status(
    item: PDFIntakeItem,
    duplicate_of: dict[str, str],
) -> str:
    """Return the initial user-facing draft status."""

    receipt = item_receipt(item)

    if item.generated_id in duplicate_of:
        return "DUPLICATE_RECEIPT"

    if item.status == "DUPLICATE_RECEIPT":
        return "DUPLICATE_RECEIPT"

    if receipt is None or not receipt.bill_found:
        return "NO_RECEIPT_FOUND"

    if item.status == "OCR_UNCLEAR" or receipt.confidence == "low":
        return "OCR_UNCLEAR"

    if not receipt.date or receipt.amount is None:
        return "EVIDENCE_DATA_MISSING"

    return "USER_INPUT_REQUIRED"


def claim_draft_row(
    sequence_number: int,
    item: PDFIntakeItem,
    duplicate_of: dict[str, str],
) -> dict[str, Any] | None:
    """Create one editable claim row from one detected receipt."""

    receipt = item_receipt(item)

    # Page-level blank/error records belong in Needs_Attention, not Claim_Draft.
    if item.candidate_no is None:
        return None

    if receipt is None or not receipt.bill_found:
        return None

    return {
        "No": sequence_number,
        "Receipt_ID": item.generated_id,
        "Month": month_label(receipt.date),
        "Extracted_Date": receipt.date,
        "Claimed_Date": receipt.date,
        "Extracted_Amount": receipt.amount,
        "Claimed_Amount": receipt.amount,
        "Currency": receipt.currency or "INR",
        "Merchant_Payee": receipt.merchant,
        "Payment_Mode": receipt.payment_mode,
        "Transaction_ID": receipt.transaction_id,
        "UTR": receipt.utr,
        "Extracted_Remark": receipt.remark,
        "Expense_Category": "",
        "Business_Purpose": receipt.remark or "",
        "Employee_ID": "",
        "Employee_Name": "",
        "Project_Cost_Centre": "",
        "User_Confirmation": "Pending",
        "Manager_Approval": "Pending",
        "Draft_Status": initial_draft_status(
            item,
            duplicate_of,
        ),
        "Source_PDF": item.source_file_name,
        "Page_No": item.page_number,
        "Candidate_No": item.candidate_no,
    }


def build_claim_draft_rows(
    results: Iterable[PDFIntakeResult],
) -> list[dict[str, Any]]:
    """Build date-sorted editable claim rows from intake results."""

    results_list = list(results)
    duplicate_of = cross_document_duplicate_map(results_list)
    rows: list[dict[str, Any]] = []

    for result in results_list:
        for item in result.items:
            row = claim_draft_row(
                sequence_number=0,
                item=item,
                duplicate_of=duplicate_of,
            )

            if row is not None:
                rows.append(row)

    rows.sort(
        key=lambda row: (
            normalise_date(row["Extracted_Date"]) or "9999-12-31",
            row["Source_PDF"].casefold(),
            row["Page_No"],
            row["Candidate_No"],
        )
    )

    month_sequences: dict[str, int] = defaultdict(int)

    for row in rows:
        month_sequences[row["Month"]] += 1
        row["No"] = month_sequences[row["Month"]]

    return rows


def build_evidence_rows(
    results: Iterable[PDFIntakeResult],
) -> list[dict[str, Any]]:
    """Build a complete machine-generated evidence register."""

    results_list = list(results)
    duplicate_of = cross_document_duplicate_map(results_list)
    rows: list[dict[str, Any]] = []

    for result in results_list:
        for item in result.items:
            rows.append(
                evidence_row(
                    result=result,
                    item=item,
                    duplicate_of=duplicate_of,
                )
            )

    return rows


def needs_attention_row(
    evidence: dict[str, Any],
) -> dict[str, Any] | None:
    """Return an issue row when evidence needs review."""

    reasons: list[str] = []

    if evidence["Intake_Status"] != "EXTRACTED_READY":
        reasons.append(str(evidence["Intake_Status"]))

    if evidence["Cross_Document_Duplicate_Of"]:
        reasons.append("CROSS_DOCUMENT_DUPLICATE")

    if not evidence["Bill_Found"]:
        reasons.append("NO_RECEIPT_FOUND")

    if evidence["Candidate_No"] is not None:
        if not evidence["Extracted_Date"]:
            reasons.append("DATE_MISSING")

        if evidence["Extracted_Amount"] is None:
            reasons.append("AMOUNT_MISSING")

        if evidence["Extraction_Confidence"] in {
            None,
            "low",
        }:
            reasons.append("CONFIDENCE_LOW")

    if not reasons:
        return None

    return {
        "Receipt_ID": evidence["Receipt_ID"],
        "Source_PDF": evidence["Source_PDF"],
        "Page_No": evidence["Page_No"],
        "Candidate_No": evidence["Candidate_No"],
        "Issue": " | ".join(dict.fromkeys(reasons)),
        "System_Message": evidence["System_Message"],
        "Evidence_Image_Path": evidence["Evidence_Image_Path"],
        "Candidate_Map_Path": evidence["Candidate_Map_Path"],
    }


def build_needs_attention_rows(
    evidence_rows: Iterable[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Build the exception/review queue."""

    rows: list[dict[str, Any]] = []

    for evidence in evidence_rows:
        issue = needs_attention_row(evidence)

        if issue is not None:
            rows.append(issue)

    return rows


def write_rows(
    worksheet,
    headers: list[str],
    rows: Iterable[dict[str, Any]],
) -> None:
    """Write a header and ordered dictionary rows to a worksheet."""

    worksheet.append(headers)

    for row in rows:
        worksheet.append(
            [row.get(header) for header in headers]
        )


def style_header(worksheet) -> None:
    """Apply consistent header formatting."""

    for cell in worksheet[1]:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=HEADER_FILL,
        )
        cell.font = Font(
            bold=True,
            color=HEADER_FONT,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def auto_width(worksheet, maximum: int = 42) -> None:
    """Set useful bounded column widths."""

    for column_cells in worksheet.columns:
        maximum_length = max(
            len(str(cell.value))
            if cell.value is not None
            else 0
            for cell in column_cells
        )

        worksheet.column_dimensions[
            column_cells[0].column_letter
        ].width = min(
            max(maximum_length + 2, 11),
            maximum,
        )


def style_standard_sheet(worksheet) -> None:
    """Apply shared table formatting."""

    worksheet.freeze_panes = "A2"
    style_header(worksheet)

    if worksheet.max_row >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    for row in worksheet.iter_rows(
        min_row=2,
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    auto_width(worksheet)


def apply_claim_draft_formatting(worksheet) -> None:
    """Colour-code editable and machine-generated claim fields."""

    style_standard_sheet(worksheet)
    header_map = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    editable_columns = {
        "Claimed_Date",
        "Claimed_Amount",
        "Expense_Category",
        "Business_Purpose",
        "Employee_ID",
        "Employee_Name",
        "Project_Cost_Centre",
        "User_Confirmation",
        "Manager_Approval",
    }

    for header, column_index in header_map.items():
        fill_colour = (
            EDITABLE_FILL
            if header in editable_columns
            else EXTRACTED_FILL
        )

        for cell in worksheet.iter_cols(
            min_col=column_index,
            max_col=column_index,
            min_row=2,
            max_row=max(worksheet.max_row, 2),
        ):
            for value_cell in cell:
                value_cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor=fill_colour,
                )

    date_columns = [
        "Extracted_Date",
        "Claimed_Date",
    ]

    for header in date_columns:
        column = header_map.get(header)

        if column:
            for cell in worksheet.iter_cols(
                min_col=column,
                max_col=column,
                min_row=2,
                max_row=worksheet.max_row,
            ):
                for value_cell in cell:
                    value_cell.number_format = "yyyy-mm-dd"

    amount_columns = [
        "Extracted_Amount",
        "Claimed_Amount",
    ]

    for header in amount_columns:
        column = header_map.get(header)

        if column:
            for cell in worksheet.iter_cols(
                min_col=column,
                max_col=column,
                min_row=2,
                max_row=worksheet.max_row,
            ):
                for value_cell in cell:
                    value_cell.number_format = '₹#,##0.00'

    category_validation = DataValidation(
        type="list",
        formula1='"' + ",".join(CATEGORY_VALUES) + '"',
        allow_blank=True,
    )
    confirmation_validation = DataValidation(
        type="list",
        formula1='"' + ",".join(CONFIRMATION_VALUES) + '"',
        allow_blank=False,
    )
    approval_validation = DataValidation(
        type="list",
        formula1='"' + ",".join(APPROVAL_VALUES) + '"',
        allow_blank=False,
    )

    worksheet.add_data_validation(category_validation)
    worksheet.add_data_validation(confirmation_validation)
    worksheet.add_data_validation(approval_validation)

    last_row = max(
        worksheet.max_row,
        2,
    )

    category_validation.add(
        worksheet.cell(
            row=2,
            column=header_map["Expense_Category"],
        ).coordinate
        + ":"
        + worksheet.cell(
            row=last_row,
            column=header_map["Expense_Category"],
        ).coordinate
    )
    confirmation_validation.add(
        worksheet.cell(
            row=2,
            column=header_map["User_Confirmation"],
        ).coordinate
        + ":"
        + worksheet.cell(
            row=last_row,
            column=header_map["User_Confirmation"],
        ).coordinate
    )
    approval_validation.add(
        worksheet.cell(
            row=2,
            column=header_map["Manager_Approval"],
        ).coordinate
        + ":"
        + worksheet.cell(
            row=last_row,
            column=header_map["Manager_Approval"],
        ).coordinate
    )

    status_column_letter = worksheet.cell(
        row=1,
        column=header_map["Draft_Status"],
    ).column_letter

    worksheet.conditional_formatting.add(
        f"A2:X{last_row}",
        FormulaRule(
            formula=[
                f'=${status_column_letter}2="DUPLICATE_RECEIPT"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor=ERROR_FILL,
            ),
        ),
    )

    worksheet.conditional_formatting.add(
        f"A2:X{last_row}",
        FormulaRule(
            formula=[
                f'=${status_column_letter}2="OCR_UNCLEAR"'
            ],
            fill=PatternFill(
                fill_type="solid",
                fgColor=WARNING_FILL,
            ),
        ),
    )

    worksheet.sheet_view.showGridLines = False


def protect_evidence_sheet(worksheet) -> None:
    """Protect evidence values from accidental editing."""

    style_standard_sheet(worksheet)

    for row in worksheet.iter_rows():
        for cell in row:
            cell.protection = Protection(
                locked=True,
            )

    worksheet.protection.sheet = True
    worksheet.protection.autoFilter = False
    worksheet.protection.sort = False
    worksheet.sheet_view.showGridLines = False


def add_instructions_sheet(workbook: Workbook) -> None:
    """Add a short, explicit user handover sheet."""

    worksheet = workbook.create_sheet(
        "Instructions",
        0,
    )

    rows = [
        [
            "PDF-FIRST REIMBURSEMENT CLAIM DRAFT",
            "",
        ],
        [
            "Purpose",
            (
                "The workbook was generated from receipt evidence. "
                "It is a draft, not an approved reimbursement."
            ),
        ],
        [
            "Yellow columns",
            "User-editable claim/business information.",
        ],
        [
            "Blue columns",
            "Machine-extracted evidence. Do not treat it as user confirmation.",
        ],
        [
            "Required user work",
            (
                "Complete Employee_Name, Expense_Category, Business_Purpose, "
                "confirm date/amount, and set User_Confirmation=Confirmed."
            ),
        ],
        [
            "Manager work",
            "Set Manager_Approval to Approved or Rejected.",
        ],
        [
            "Validation",
            (
                "Upload the completed workbook to the Validate Completed Draft "
                "tab. Validation compares user values with locked evidence."
            ),
        ],
        [
            "Safety",
            (
                "Document intake and claim-draft validation do not transfer "
                "money and do not directly approve payment."
            ),
        ],
    ]

    for row in rows:
        worksheet.append(row)

    worksheet["A1"].font = Font(
        bold=True,
        size=16,
        color=HEADER_FONT,
    )
    worksheet["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor=HEADER_FILL,
    )
    worksheet["B1"].fill = PatternFill(
        fill_type="solid",
        fgColor=HEADER_FILL,
    )

    for row in worksheet.iter_rows(
        min_row=2,
    ):
        row[0].font = Font(
            bold=True,
        )
        row[0].fill = PatternFill(
            fill_type="solid",
            fgColor=EXTRACTED_FILL,
        )
        row[1].alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 90
    worksheet.sheet_view.showGridLines = False


def build_summary_rows(
    results: list[PDFIntakeResult],
    claim_rows: list[dict[str, Any]],
    evidence_rows: list[dict[str, Any]],
    attention_rows: list[dict[str, Any]],
) -> list[list[Any]]:
    """Build summary metrics and monthly totals."""

    total_amount = sum(
        (
            normalise_amount(row["Extracted_Amount"])
            or Decimal("0.00")
        )
        for row in claim_rows
        if row["Draft_Status"] != "DUPLICATE_RECEIPT"
    )

    rows: list[list[Any]] = [
        ["Metric", "Value"],
        ["Input documents", len(results)],
        [
            "Total pages",
            sum(result.total_pages for result in results),
        ],
        [
            "Detected candidates",
            sum(result.total_candidates for result in results),
        ],
        ["Claim draft rows", len(claim_rows)],
        ["Evidence records", len(evidence_rows)],
        ["Needs attention", len(attention_rows)],
        [
            "Duplicate records",
            sum(
                row["Draft_Status"] == "DUPLICATE_RECEIPT"
                for row in claim_rows
            ),
        ],
        ["Extracted amount total", float(total_amount)],
        [],
        ["Monthly extracted totals", ""],
        ["Month", "Amount"],
    ]

    monthly_totals: dict[str, Decimal] = defaultdict(
        lambda: Decimal("0.00")
    )

    for row in claim_rows:
        if row["Draft_Status"] == "DUPLICATE_RECEIPT":
            continue

        amount = normalise_amount(
            row["Extracted_Amount"]
        )

        if amount is not None:
            monthly_totals[row["Month"]] += amount

    for month, amount in sorted(
        monthly_totals.items()
    ):
        rows.append(
            [
                month,
                float(amount),
            ]
        )

    return rows


def write_claim_draft_workbook(
    results: Iterable[PDFIntakeResult],
    output_path: str | Path,
) -> Path:
    """Create the complete PDF-first reimbursement draft workbook."""

    results_list = list(results)
    output = Path(output_path)
    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    claim_rows = build_claim_draft_rows(
        results_list
    )
    evidence_rows = build_evidence_rows(
        results_list
    )
    attention_rows = build_needs_attention_rows(
        evidence_rows
    )

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    add_instructions_sheet(workbook)

    summary_sheet = workbook.create_sheet(
        "Summary"
    )
    for row in build_summary_rows(
        results=results_list,
        claim_rows=claim_rows,
        evidence_rows=evidence_rows,
        attention_rows=attention_rows,
    ):
        summary_sheet.append(row)

    style_standard_sheet(summary_sheet)
    summary_sheet.column_dimensions["A"].width = 30
    summary_sheet.column_dimensions["B"].width = 35

    for cell in summary_sheet["B"]:
        if isinstance(cell.value, (float, int)):
            cell.number_format = '₹#,##0.00'

    claim_sheet = workbook.create_sheet(
        "Claim_Draft"
    )
    write_rows(
        worksheet=claim_sheet,
        headers=CLAIM_DRAFT_HEADERS,
        rows=claim_rows,
    )
    apply_claim_draft_formatting(
        claim_sheet
    )

    evidence_sheet = workbook.create_sheet(
        "Evidence_Register"
    )
    write_rows(
        worksheet=evidence_sheet,
        headers=EVIDENCE_HEADERS,
        rows=evidence_rows,
    )
    protect_evidence_sheet(
        evidence_sheet
    )

    attention_sheet = workbook.create_sheet(
        "Needs_Attention"
    )
    attention_headers = [
        "Receipt_ID",
        "Source_PDF",
        "Page_No",
        "Candidate_No",
        "Issue",
        "System_Message",
        "Evidence_Image_Path",
        "Candidate_Map_Path",
    ]
    write_rows(
        worksheet=attention_sheet,
        headers=attention_headers,
        rows=attention_rows,
    )
    style_standard_sheet(
        attention_sheet
    )

    validation_sheet = workbook.create_sheet(
        "Validation_Result"
    )
    write_rows(
        worksheet=validation_sheet,
        headers=VALIDATION_HEADERS,
        rows=[],
    )
    style_standard_sheet(
        validation_sheet
    )

    input_sheet = workbook.create_sheet(
        "Input_Documents"
    )
    input_rows = [
        result.ingestion.model_dump(
            mode="json"
        )
        for result in results_list
    ]
    input_headers = sorted(
        {
            key
            for row in input_rows
            for key in row
        }
    )
    write_rows(
        worksheet=input_sheet,
        headers=input_headers,
        rows=input_rows,
    )
    protect_evidence_sheet(
        input_sheet
    )

    workbook.save(
        output
    )

    return output


def worksheet_rows_by_header(
    worksheet,
) -> tuple[
    dict[str, int],
    list[dict[str, Any]],
]:
    """Read a worksheet into header-indexed dictionaries."""

    headers = {
        str(cell.value).strip(): cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }

    rows: list[dict[str, Any]] = []

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        values = {
            header: worksheet.cell(
                row=row_number,
                column=column,
            ).value
            for header, column in headers.items()
        }

        if not any(
            value not in (None, "")
            for value in values.values()
        ):
            continue

        values["_row_number"] = row_number
        rows.append(values)

    return headers, rows


def validation_status(
    claim: dict[str, Any],
    evidence: dict[str, Any] | None,
) -> tuple[str, str]:
    """Validate one user-completed claim against locked evidence."""

    if evidence is None:
        return (
            "EVIDENCE_REFERENCE_MISSING",
            "Receipt_ID was not found in Evidence_Register.",
        )

    if evidence.get(
        "Cross_Document_Duplicate_Of"
    ) or evidence.get("Duplicate_Of"):
        return (
            "DUPLICATE_RECEIPT",
            "The evidence is marked as a duplicate receipt.",
        )

    intake_status = evidence.get(
        "Intake_Status"
    )

    if intake_status in {
        "OCR_UNCLEAR",
        "NO_RECEIPT_FOUND",
        "EMPTY_PAGE",
        "EXTRACTION_FAILED",
    }:
        return (
            str(intake_status),
            "The evidence is not suitable for claim submission.",
        )

    confirmation = str(
        claim.get("User_Confirmation")
        or ""
    ).strip()

    if confirmation.casefold() == "rejected":
        return (
            "USER_REJECTED",
            "The employee rejected this draft claim row.",
        )

    if confirmation.casefold() != "confirmed":
        return (
            "USER_CONFIRMATION_PENDING",
            "Set User_Confirmation to Confirmed after checking the row.",
        )

    required_fields = {
        "Employee_Name": claim.get(
            "Employee_Name"
        ),
        "Expense_Category": claim.get(
            "Expense_Category"
        ),
        "Business_Purpose": claim.get(
            "Business_Purpose"
        ),
        "Claimed_Date": claim.get(
            "Claimed_Date"
        ),
        "Claimed_Amount": claim.get(
            "Claimed_Amount"
        ),
    }

    missing_fields = [
        field
        for field, value in required_fields.items()
        if value is None
        or not str(value).strip()
    ]

    if missing_fields:
        return (
            "REQUIRED_BUSINESS_DATA_MISSING",
            "Missing required claim fields: "
            + ", ".join(missing_fields)
            + ".",
        )

    evidence_date = normalise_date(
        evidence.get("Extracted_Date")
    )
    claimed_date = normalise_date(
        claim.get("Claimed_Date")
    )

    if not evidence_date:
        return (
            "EVIDENCE_DATE_MISSING",
            "The receipt evidence does not contain a reliable date.",
        )

    if not claimed_date:
        return (
            "CLAIM_DATE_INVALID",
            "Claimed_Date could not be parsed.",
        )

    if evidence_date != claimed_date:
        return (
            "DATE_CHANGED",
            (
                f"Claimed date {claimed_date} differs from "
                f"evidence date {evidence_date}."
            ),
        )

    evidence_amount = normalise_amount(
        evidence.get("Extracted_Amount")
    )
    claimed_amount = normalise_amount(
        claim.get("Claimed_Amount")
    )

    if evidence_amount is None:
        return (
            "EVIDENCE_AMOUNT_MISSING",
            "The receipt evidence does not contain a reliable amount.",
        )

    if claimed_amount is None:
        return (
            "CLAIM_AMOUNT_INVALID",
            "Claimed_Amount could not be parsed.",
        )

    if evidence_amount != claimed_amount:
        return (
            "AMOUNT_CHANGED",
            (
                f"Claimed amount {claimed_amount} differs from "
                f"evidence amount {evidence_amount}."
            ),
        )

    manager_approval = str(
        claim.get("Manager_Approval")
        or ""
    ).strip()

    if manager_approval.casefold() == "rejected":
        return (
            "MANAGER_REJECTED",
            "The manager rejected this reimbursement claim.",
        )

    if manager_approval.casefold() != "approved":
        return (
            "MANAGER_APPROVAL_PENDING",
            "The claim evidence matches, but manager approval is pending.",
        )

    return (
        "OK_VERIFIED",
        (
            "User-confirmed claim values match the locked receipt evidence "
            "and manager approval is recorded."
        ),
    )


def validate_completed_claim_draft(
    input_workbook: str | Path,
    output_path: str | Path | None = None,
) -> tuple[Path, list[dict[str, Any]]]:
    """Validate an edited Claim_Draft sheet against Evidence_Register."""

    source = Path(input_workbook)

    if output_path is None:
        output = source.with_name(
            source.stem
            + "_validated.xlsx"
        )
    else:
        output = Path(output_path)

    workbook = load_workbook(
        source
    )

    required_sheets = {
        "Claim_Draft",
        "Evidence_Register",
    }

    missing_sheets = required_sheets.difference(
        workbook.sheetnames
    )

    if missing_sheets:
        raise ValueError(
            "Missing required workbook sheets: "
            + ", ".join(
                sorted(missing_sheets)
            )
        )

    claim_sheet = workbook[
        "Claim_Draft"
    ]
    evidence_sheet = workbook[
        "Evidence_Register"
    ]

    claim_headers, claim_rows = worksheet_rows_by_header(
        claim_sheet
    )
    _, evidence_rows = worksheet_rows_by_header(
        evidence_sheet
    )

    required_claim_columns = {
        "Receipt_ID",
        "Claimed_Date",
        "Claimed_Amount",
        "Expense_Category",
        "Business_Purpose",
        "Employee_Name",
        "User_Confirmation",
        "Manager_Approval",
    }

    missing_columns = required_claim_columns.difference(
        claim_headers
    )

    if missing_columns:
        raise ValueError(
            "Claim_Draft is missing required columns: "
            + ", ".join(
                sorted(missing_columns)
            )
        )

    evidence_by_id = {
        str(row.get("Receipt_ID")).strip(): row
        for row in evidence_rows
        if row.get("Receipt_ID")
    }

    validation_rows: list[
        dict[str, Any]
    ] = []

    seen_receipt_ids: set[str] = set()

    for claim in claim_rows:
        receipt_id = str(
            claim.get("Receipt_ID")
            or ""
        ).strip()
        evidence = evidence_by_id.get(
            receipt_id
        )

        if receipt_id in seen_receipt_ids:
            status = "DUPLICATE_CLAIM_ROW"
            reason = (
                "Receipt_ID appears more than once in Claim_Draft."
            )
        else:
            status, reason = validation_status(
                claim=claim,
                evidence=evidence,
            )

        seen_receipt_ids.add(
            receipt_id
        )

        validation_rows.append(
            {
                "Receipt_ID": receipt_id,
                "Employee_ID": claim.get(
                    "Employee_ID"
                ),
                "Employee_Name": claim.get(
                    "Employee_Name"
                ),
                "Expense_Category": claim.get(
                    "Expense_Category"
                ),
                "Business_Purpose": claim.get(
                    "Business_Purpose"
                ),
                "Evidence_Date": (
                    evidence.get(
                        "Extracted_Date"
                    )
                    if evidence
                    else None
                ),
                "Claimed_Date": claim.get(
                    "Claimed_Date"
                ),
                "Evidence_Amount": (
                    evidence.get(
                        "Extracted_Amount"
                    )
                    if evidence
                    else None
                ),
                "Claimed_Amount": claim.get(
                    "Claimed_Amount"
                ),
                "User_Confirmation": claim.get(
                    "User_Confirmation"
                ),
                "Manager_Approval": claim.get(
                    "Manager_Approval"
                ),
                "Validation_Status": status,
                "Validation_Reason": reason,
                "Source_PDF": (
                    evidence.get(
                        "Source_PDF"
                    )
                    if evidence
                    else None
                ),
                "Page_No": (
                    evidence.get(
                        "Page_No"
                    )
                    if evidence
                    else None
                ),
                "Candidate_No": (
                    evidence.get(
                        "Candidate_No"
                    )
                    if evidence
                    else None
                ),
            }
        )

    if "Validation_Result" in workbook.sheetnames:
        index = workbook.sheetnames.index(
            "Validation_Result"
        )
        workbook.remove(
            workbook[
                "Validation_Result"
            ]
        )
        validation_sheet = workbook.create_sheet(
            "Validation_Result",
            index,
        )
    else:
        validation_sheet = workbook.create_sheet(
            "Validation_Result"
        )

    write_rows(
        worksheet=validation_sheet,
        headers=VALIDATION_HEADERS,
        rows=validation_rows,
    )
    style_standard_sheet(
        validation_sheet
    )

    validation_headers = {
        cell.value: cell.column
        for cell in validation_sheet[1]
    }
    status_column = validation_headers[
        "Validation_Status"
    ]

    for row_number in range(
        2,
        validation_sheet.max_row + 1,
    ):
        status = str(
            validation_sheet.cell(
                row=row_number,
                column=status_column,
            ).value
            or ""
        )

        if status == "OK_VERIFIED":
            fill = GOOD_FILL
        elif status in {
            "USER_CONFIRMATION_PENDING",
            "MANAGER_APPROVAL_PENDING",
        }:
            fill = WARNING_FILL
        else:
            fill = ERROR_FILL

        for cell in validation_sheet[
            row_number
        ]:
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=fill,
            )

    if "Validation_Status" not in claim_headers:
        validation_status_column = (
            claim_sheet.max_column + 1
        )
        claim_sheet.cell(
            row=1,
            column=validation_status_column,
            value="Validation_Status",
        )
        claim_sheet.cell(
            row=1,
            column=validation_status_column,
        ).fill = PatternFill(
            fill_type="solid",
            fgColor=HEADER_FILL,
        )
        claim_sheet.cell(
            row=1,
            column=validation_status_column,
        ).font = Font(
            bold=True,
            color=HEADER_FONT,
        )
    else:
        validation_status_column = (
            claim_headers[
                "Validation_Status"
            ]
        )

    status_by_id = {
        row["Receipt_ID"]: row[
            "Validation_Status"
        ]
        for row in validation_rows
    }

    receipt_id_column = claim_headers[
        "Receipt_ID"
    ]

    for row_number in range(
        2,
        claim_sheet.max_row + 1,
    ):
        receipt_id = str(
            claim_sheet.cell(
                row=row_number,
                column=receipt_id_column,
            ).value
            or ""
        ).strip()

        claim_sheet.cell(
            row=row_number,
            column=validation_status_column,
            value=status_by_id.get(
                receipt_id,
                "NOT_VALIDATED",
            ),
        )

    auto_width(
        claim_sheet
    )
    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )
    workbook.save(
        output
    )

    return output, validation_rows
