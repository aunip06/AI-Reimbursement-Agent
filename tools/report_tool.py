import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models.worksheet_processing_result import (
    WorksheetProcessingResult,
)


RESULT_COLUMNS = [
    "Employee_ID",
    "Employee_Name",
    "Expense_ID",
    "Excel_Row",
    "Excel_Date",
    "Expense_Category",
    "Description",
    "Excel_Amount",
    "Excel_Vendor",
    "Excel_Payment_Mode",
    "Excel_Transaction_ID",
    "Excel_Invoice_No",
    "Monthly_PDF_Name",
    "Receipt_Page_No",
    "Result_Source",
    "Final_Status",
    "Auto_Payable",
    "Message",
    "Extracted_Bill_Found",
    "Extracted_Receipt_Type",
    "Extracted_Date",
    "Extracted_Amount",
    "Extracted_Merchant",
    "Extracted_Payment_Mode",
    "Extracted_Transaction_ID",
    "Extracted_UTR",
    "Extracted_Invoice_No",
    "Receipt_Confidence",
    "Analysis_Confidence",
    "Agent_Recommended_Status",
    "Agent_Reasons",
    "Final_Reasons",
    "Safety_Checks",
]


EVIDENCE_COLUMNS = [
    "Expense_ID",
    "PDF_File_Name",
    "Page_Number",
    "Evidence_Status",
    "Image_Path",
    "Direct_PDF_Text",
    "Local_OCR_Text",
    "Combined_Evidence_Text",
]


def get_processing_objects(
    worksheet_claim_result,
):
    """
    Safely obtain nested processing objects.
    """

    workflow_result = (
        worksheet_claim_result.workflow_result
    )

    if workflow_result is None:
        return None, None, None

    processing_result = (
        workflow_result.processing_result
    )

    if processing_result is None:
        return None, None, None

    analysis = processing_result.analysis
    final_decision = processing_result.final_decision

    return processing_result, analysis, final_decision


def build_result_rows(
    worksheet_result: WorksheetProcessingResult,
) -> list[dict[str, Any]]:
    """
    Flatten worksheet results into report rows.
    """

    rows: list[dict[str, Any]] = []

    for item in worksheet_result.results:
        claim = item.claim

        (
            processing_result,
            analysis,
            final_decision,
        ) = get_processing_objects(item)

        receipt = (
            analysis.receipt
            if analysis is not None
            else None
        )

        agent_reasons = (
            " | ".join(analysis.reasons)
            if analysis is not None
            else ""
        )

        final_reasons = (
            " | ".join(final_decision.reasons)
            if final_decision is not None
            else item.message
        )

        safety_checks = (
            json.dumps(
                final_decision.safety_checks,
                ensure_ascii=False,
            )
            if final_decision is not None
            else ""
        )

        rows.append(
            {
                "Employee_ID": claim.employee_id,
                "Employee_Name": claim.employee_name,
                "Expense_ID": claim.expense_id,
                "Excel_Row": claim.source_row_number,
                "Excel_Date": claim.expense_date.isoformat(),
                "Expense_Category": claim.expense_category,
                "Description": claim.description,
                "Excel_Amount": float(
                    claim.claimed_amount
                ),
                "Excel_Vendor": claim.vendor,
                "Excel_Payment_Mode": claim.payment_mode,
                "Excel_Transaction_ID": (
                    claim.transaction_id
                ),
                "Excel_Invoice_No": claim.invoice_no,
                "Monthly_PDF_Name": (
                    claim.monthly_pdf_name
                ),
                "Receipt_Page_No": (
                    claim.receipt_page_no
                ),
                "Result_Source": item.result_source,
                "Final_Status": item.final_status,
                "Auto_Payable": item.auto_payable,
                "Message": item.message,
                "Extracted_Bill_Found": (
                    receipt.bill_found
                    if receipt is not None
                    else None
                ),
                "Extracted_Receipt_Type": (
                    receipt.receipt_type
                    if receipt is not None
                    else None
                ),
                "Extracted_Date": (
                    receipt.date
                    if receipt is not None
                    else None
                ),
                "Extracted_Amount": (
                    receipt.amount
                    if receipt is not None
                    else None
                ),
                "Extracted_Merchant": (
                    receipt.merchant
                    if receipt is not None
                    else None
                ),
                "Extracted_Payment_Mode": (
                    receipt.payment_mode
                    if receipt is not None
                    else None
                ),
                "Extracted_Transaction_ID": (
                    receipt.transaction_id
                    if receipt is not None
                    else None
                ),
                "Extracted_UTR": (
                    receipt.utr
                    if receipt is not None
                    else None
                ),
                "Extracted_Invoice_No": (
                    receipt.invoice_no
                    if receipt is not None
                    else None
                ),
                "Receipt_Confidence": (
                    receipt.confidence
                    if receipt is not None
                    else None
                ),
                "Analysis_Confidence": (
                    analysis.analysis_confidence
                    if analysis is not None
                    else None
                ),
                "Agent_Recommended_Status": (
                    analysis.recommended_status
                    if analysis is not None
                    else None
                ),
                "Agent_Reasons": agent_reasons,
                "Final_Reasons": final_reasons,
                "Safety_Checks": safety_checks,
            }
        )

    return rows


def build_evidence_rows(
    worksheet_result: WorksheetProcessingResult,
) -> list[dict[str, Any]]:
    """
    Create page-level OCR and PDF evidence rows.
    """

    rows: list[dict[str, Any]] = []

    for item in worksheet_result.results:
        workflow_result = item.workflow_result

        if workflow_result is None:
            continue

        evidence = workflow_result.evidence

        rows.append(
            {
                "Expense_ID": item.claim.expense_id,
                "PDF_File_Name": evidence.pdf_file_name,
                "Page_Number": evidence.page_number,
                "Evidence_Status": evidence.status,
                "Image_Path": evidence.image_path,
                "Direct_PDF_Text": (
                    evidence.direct_pdf_text
                ),
                "Local_OCR_Text": evidence.ocr_text,
                "Combined_Evidence_Text": (
                    evidence.combined_evidence_text
                ),
            }
        )

    return rows


def create_summary_dataframe(
    worksheet_result: WorksheetProcessingResult,
    result_rows: list[dict[str, Any]],
) -> pd.DataFrame:
    """
    Create worksheet-level report metrics.
    """

    approved_amount = sum(
        float(row["Excel_Amount"])
        for row in result_rows
        if row["Auto_Payable"] is True
    )

    openai_count = sum(
        row["Result_Source"] == "openai"
        for row in result_rows
    )

    cache_count = sum(
        row["Result_Source"] == "cache"
        for row in result_rows
    )

    local_count = sum(
        row["Result_Source"] == "local_validation"
        for row in result_rows
    )

    summary_rows = [
        {
            "Metric": "Worksheet",
            "Value": worksheet_result.worksheet_name,
        },
        {
            "Metric": "Total valid claims",
            "Value": worksheet_result.total_claims,
        },
        {
            "Metric": "Completed SDK/cache claims",
            "Value": worksheet_result.completed_claims,
        },
        {
            "Metric": "Approved claims",
            "Value": worksheet_result.approved_count,
        },
        {
            "Metric": "Exception claims",
            "Value": worksheet_result.exception_count,
        },
        {
            "Metric": "Duplicate-page claims",
            "Value": (
                worksheet_result.duplicate_page_claims
            ),
        },
        {
            "Metric": "OpenAI results",
            "Value": openai_count,
        },
        {
            "Metric": "Cached results",
            "Value": cache_count,
        },
        {
            "Metric": "Local validation results",
            "Value": local_count,
        },
        {
            "Metric": "Excel format errors",
            "Value": len(
                worksheet_result.excel_errors
            ),
        },
        {
            "Metric": "Approved amount",
            "Value": round(
                approved_amount,
                2,
            ),
        },
    ]

    return pd.DataFrame(summary_rows)


def style_report_workbook(
    workbook,
) -> None:
    """
    Apply readable formatting to every generated worksheet.
    """

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    approved_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )

    exception_fill = PatternFill(
        fill_type="solid",
        fgColor="FCE4D6",
    )

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for row in worksheet.iter_rows(
            min_row=2
        ):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        for column_index, column_cells in enumerate(
            worksheet.columns,
            start=1,
        ):
            maximum_length = 0

            for cell in column_cells:
                value = cell.value

                if value is None:
                    continue

                maximum_length = max(
                    maximum_length,
                    len(str(value)),
                )

            worksheet.column_dimensions[
                get_column_letter(column_index)
            ].width = min(
                max(maximum_length + 2, 12),
                40,
            )

        if (
            worksheet.title
            == "Approved_For_Accounts"
        ):
            for row in worksheet.iter_rows(
                min_row=2
            ):
                for cell in row:
                    cell.fill = approved_fill

        if worksheet.title == "Exception_Report":
            for row in worksheet.iter_rows(
                min_row=2
            ):
                for cell in row:
                    cell.fill = exception_fill

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )


def generate_worksheet_report(
    worksheet_result: WorksheetProcessingResult,
    output_path: str | Path,
) -> Path:
    """
    Generate the final monthly reimbursement Excel report.
    """

    report_path = Path(output_path)

    report_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    result_rows = build_result_rows(
        worksheet_result
    )

    evidence_rows = build_evidence_rows(
        worksheet_result
    )

    all_results_df = pd.DataFrame(
        result_rows,
        columns=RESULT_COLUMNS,
    )

    if all_results_df.empty:
        approved_df = pd.DataFrame(
            columns=RESULT_COLUMNS
        )

        exception_df = pd.DataFrame(
            columns=RESULT_COLUMNS
        )

    else:
        approved_df = all_results_df[
            all_results_df["Auto_Payable"]
            == True
        ].copy()

        exception_df = all_results_df[
            all_results_df["Auto_Payable"]
            == False
        ].copy()

    evidence_df = pd.DataFrame(
        evidence_rows,
        columns=EVIDENCE_COLUMNS,
    )

    excel_errors_df = pd.DataFrame(
        worksheet_result.excel_errors
    )

    summary_df = create_summary_dataframe(
        worksheet_result=worksheet_result,
        result_rows=result_rows,
    )

    with pd.ExcelWriter(
        report_path,
        engine="openpyxl",
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="Summary",
            index=False,
        )

        approved_df.to_excel(
            writer,
            sheet_name="Approved_For_Accounts",
            index=False,
        )

        exception_df.to_excel(
            writer,
            sheet_name="Exception_Report",
            index=False,
        )

        all_results_df.to_excel(
            writer,
            sheet_name="All_Results",
            index=False,
        )

        evidence_df.to_excel(
            writer,
            sheet_name="AI_Page_Evidence",
            index=False,
        )

        excel_errors_df.to_excel(
            writer,
            sheet_name="Excel_Errors",
            index=False,
        )

        style_report_workbook(
            writer.book
        )

    return report_path