from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from models.pdf_intake import PDFIntakeItem, PDFIntakeResult


def item_to_row(item: PDFIntakeItem) -> dict:
    receipt = item.receipt

    return {
        "Generated_ID": item.generated_id,
        "Source_File": item.source_file_name,
        "Page_Number": item.page_number,
        "Candidate_Number": item.candidate_no,
        "Intake_Status": item.status,
        "Result_Source": item.result_source,
        "Bill_Found": receipt.bill_found if receipt else None,
        "Receipt_Type": receipt.receipt_type if receipt else None,
        "Date": receipt.date if receipt else None,
        "Amount": receipt.amount if receipt else None,
        "Currency": receipt.currency if receipt else None,
        "Remark": getattr(receipt, "remark", None) if receipt else None,
        "Merchant": receipt.merchant if receipt else None,
        "Payment_Mode": receipt.payment_mode if receipt else None,
        "Payment_App": receipt.payment_app if receipt else None,
        "Payment_Message": receipt.message if receipt else None,
        "Transaction_ID": receipt.transaction_id if receipt else None,
        "UTR": receipt.utr if receipt else None,
        "Invoice_No": receipt.invoice_no if receipt else None,
        "Receipt_No": receipt.receipt_no if receipt else None,
        "GST_Number": receipt.gst_number if receipt else None,
        "Extraction_Confidence": receipt.confidence if receipt else None,
        "Extraction_Notes": receipt.extraction_notes if receipt else None,
        "Duplicate_Of": item.duplicate_of_generated_id,
        "Candidate_Fingerprint": item.candidate_fingerprint,
        "Perceptual_Hash": item.perceptual_hash,
        "Evidence_Image_Path": item.candidate_image_path,
        "Candidate_Map_Path": item.candidate_map_image_path,
        "Message": item.message,
    }


def build_intake_rows(result: PDFIntakeResult) -> list[dict]:
    return [item_to_row(item) for item in result.items]


def style_worksheet(worksheet) -> None:
    worksheet.freeze_panes = "A2"

    if worksheet.max_row >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="FF1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for column_cells in worksheet.columns:
        maximum_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )

        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(maximum_length + 2, 12),
            50,
        )


def write_pdf_intake_report(
    result: PDFIntakeResult,
    report_path: str | Path,
) -> Path:
    """Write one auditable document-intake workbook."""

    output_path = Path(report_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    all_dataframe = pd.DataFrame(build_intake_rows(result))

    if all_dataframe.empty:
        extracted_dataframe = all_dataframe.copy()
        exception_dataframe = all_dataframe.copy()
    else:
        extracted_dataframe = all_dataframe[
            all_dataframe["Intake_Status"] == "EXTRACTED_READY"
        ].copy()

        exception_dataframe = all_dataframe[
            all_dataframe["Intake_Status"] != "EXTRACTED_READY"
        ].copy()

    summary_dataframe = pd.DataFrame(
        [
            {"Metric": "Source file", "Value": result.ingestion.source_file_name},
            {"Metric": "Input family", "Value": result.ingestion.source_family},
            {"Metric": "Ingestion status", "Value": result.ingestion.status},
            {"Metric": "Converter", "Value": result.ingestion.converter},
            {"Metric": "Total pages", "Value": result.total_pages},
            {"Metric": "Detected candidates", "Value": result.total_candidates},
            {"Metric": "Extracted ready", "Value": result.extracted_count},
            {"Metric": "Exceptions", "Value": result.exception_count},
            {"Metric": "Duplicate receipts", "Value": result.duplicate_count},
            {"Metric": "Result message", "Value": result.message},
        ]
    )

    ingestion_dataframe = pd.DataFrame(
        [result.ingestion.model_dump(mode="json")]
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_dataframe.to_excel(writer, sheet_name="Summary", index=False)
        extracted_dataframe.to_excel(
            writer,
            sheet_name="Extracted_Receipts",
            index=False,
        )
        exception_dataframe.to_excel(
            writer,
            sheet_name="Exceptions",
            index=False,
        )
        all_dataframe.to_excel(writer, sheet_name="All_Results", index=False)
        ingestion_dataframe.to_excel(
            writer,
            sheet_name="Input_Document",
            index=False,
        )

        for worksheet in writer.book.worksheets:
            style_worksheet(worksheet)

    return output_path
